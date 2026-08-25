"""
Reusable multiple-choice marker for scanned OMR answer sheets.

Typical use (same school answer-sheet template, new test):
    1. Put the scanned PDF and the model-answer .docx in this folder.
    2. Edit the CONFIG block below (PDF_PATH, KEY, OUTPUT, OVERRIDES if any).
    3. Run:  python mark_mc.py
       -> produces the results .xlsx

If you switch to a DIFFERENT printed answer sheet, the bubble positions change.
Run:  python mark_mc.py --calibrate
This saves "calibration.png": page 1 with the sampling points drawn on top.
Adjust GEOMETRY until every dot sits inside its A/B/C/D box, then re-run.

Marking rules:
  - Each bubble is read by how dark it is. The darkest box per question wins.
  - Blank / faint / double-marked / crossed-out cells are auto-flagged in the
    "Review Qs" column so you can eyeball them.
  - Put hand-verified answers in OVERRIDES to correct any flagged cell.
"""

import sys, json, os, statistics
import numpy as np
from PIL import Image, ImageDraw

# ============================ CONFIG ============================
PDF_PATH = "F5 maths Ex2 MC.pdf"
KEY      = "docx:ans.docx"          # "docx:<file>" reads the answer table, OR a literal string like "DCBBC..."
NUM_Q    = None                     # None = use however many answers the key has
OUTPUT   = "F5 Maths Ex2 MC - Results.xlsx"
DPI      = 300
PASS_MARK = 0.5                     # fraction needed to "pass" (set None to hide pass stats)

# Hand-verified corrections for flagged cells:  (page, question): "A"/"B"/"C"/"D", "-" blank, "X" invalid
OVERRIDES = {
 (38,11):'X', (38,18):'X',
 (42,17):'D', (42,18):'D',
 (45,20):'B', (47,20):'B',
 (64,7):'B', (64,8):'A', (64,9):'D', (64,19):'B', (64,20):'C', (64,21):'C',
 (69,16):'D', (79,21):'B',
 (86,9):'D', (86,21):'A', (86,22):'A',
}

# --- GEOMETRY of the printed sheet (pixel coords at DPI above) ---
# Three column-blocks, each holding 20 questions (rows 1-10 upper, 11-20 lower).
# Block 1 = Q1-20, Block 2 = Q21-40, Block 3 = Q41-60.
GEOMETRY = dict(
    blocks=[                                   # x-centre of each A/B/C/D bubble per block
        {'A':565, 'B':705, 'C':840,  'D':975},
        {'A':1265,'B':1405,'C':1545, 'D':1680},
        {'A':1965,'B':2105,'C':2245, 'D':2385},  # block 3 - VERIFY with --calibrate if you mark >40 Qs
    ],
    y_upper0=970, y_upper_step=85,             # row centres 1-10  : y = y_upper0 + (row-1)*step
    y_lower0=1875, y_lower_step=84,            # row centres 11-20 : y = y_lower0 + (row-11)*step
    half_w=40, half_h=13,                      # half-size of the sampling window
)
# ===============================================================


def load_key(spec):
    if spec.startswith("docx:"):
        from docx import Document
        doc = Document(spec[5:])
        pairs = []
        for t in doc.tables:
            for row in t.rows:
                cells = [c.text.strip() for c in row.cells]
                if len(cells) >= 2 and cells[0].isdigit() and cells[1].upper() in "ABCD":
                    pairs.append((int(cells[0]), cells[1].upper()))
        pairs.sort()
        return [p[1] for p in pairs]
    return list(spec.strip().upper())


def targets(num_q, g):
    """Return list of (question, {A,B,C,D x-centres}, base_y)."""
    out = []
    for q in range(1, num_q + 1):
        b = (q - 1) // 20
        within = (q - 1) % 20            # 0..19
        if b >= len(g['blocks']):
            raise ValueError("Question %d exceeds 3 blocks (60 Q max)" % q)
        xc = g['blocks'][b]
        if within < 10:
            y = g['y_upper0'] + within * g['y_upper_step']
        else:
            y = g['y_lower0'] + (within - 10) * g['y_lower_step']
        out.append((q, xc, y))
    return out


def darkness(a, cx, cy, g):
    b = a[cy - g['half_h']:cy + g['half_h'], cx - g['half_w']:cx + g['half_w']]
    return 0.0 if b.size == 0 else 255 - b.mean()


def global_offset(a, T, g):
    best = None
    for dy in range(-26, 27, 3):
        for dx in range(-26, 27, 3):
            tot = 0
            for q, xc, yc in T:
                ds = sorted([darkness(a, xc[k] + dx, yc + dy, g) for k in 'ABCD'], reverse=True)
                tot += ds[0] - ds[1]
            if best is None or tot > best[0]:
                best = (tot, dx, dy)
    return best[1], best[2]


def yrefine(a, xc, yc0, dx, g):
    by, bsc = yc0, -1
    for ddy in range(-8, 9):
        vals = [darkness(a, xc[k] + dx, yc0 + ddy, g) for k in 'ABCD']
        sc = max(vals) - statistics.median(vals)
        if sc > bsc:
            bsc, by = sc, yc0 + ddy
    return by


def read_page(path, T, g):
    a = np.array(Image.open(path).convert('L'))
    dx, dy = global_offset(a, T, g)
    res = {}
    for q, xc, yc in T:
        by = yrefine(a, xc, yc + dy, dx, g)
        d = {}
        for k in 'ABCD':
            bv = -1
            for ddx in range(-6, 7, 2):
                bv = max(bv, darkness(a, xc[k] + dx + ddx, by, g))
            d[k] = bv
        o = sorted(d, key=d.get, reverse=True)
        v = sorted(d.values(), reverse=True)
        d0, d1, d3 = v[0], v[1], v[3]
        med = statistics.median(d.values())
        spread = d0 - d3
        ans, flag = o[0], ''
        if d0 < 48 or (d0 - med) < 11:
            ans, flag = '', 'BLANK'
        elif (d0 - d1) < 0.22 * spread and (d0 - d1) < 13:
            flag = 'AMBIG'
        elif d1 >= 48 and (d1 - med) >= 11 and (d0 - d1) < 0.5 * spread:
            flag = 'MULTI'
        res[q] = {'ans': ans, 'd': {k: round(d[k], 1) for k in 'ABCD'}, 'flag': flag}
    return res


def render_pages(pdf, dpi, outdir):
    import fitz
    os.makedirs(outdir, exist_ok=True)
    doc = fitz.open(pdf)
    paths = []
    for i in range(doc.page_count):
        p = os.path.join(outdir, "p_%03d.png" % (i + 1))
        doc[i].get_pixmap(dpi=dpi).save(p)
        paths.append(p)
    return paths


def calibrate(pdf, dpi, num_q, g):
    import fitz
    doc = fitz.open(pdf)
    pix = doc[0].get_pixmap(dpi=dpi)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples).convert("RGB")
    dr = ImageDraw.Draw(img)
    for q, xc, yc in targets(num_q, g):
        for k in 'ABCD':
            x, y = xc[k], yc
            dr.ellipse([x - 5, y - 5, x + 5, y + 5], outline=(255, 0, 0), width=2)
            dr.rectangle([x - g['half_w'], y - g['half_h'], x + g['half_w'], y + g['half_h']],
                         outline=(0, 120, 255), width=1)
    img.save("calibration.png")
    print("Wrote calibration.png - red dots must sit inside the A/B/C/D boxes for every question.")


def build_excel(pages, ans, flags, key, out, pass_mark):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    nq = len(key); n = len(pages); first = 3; last = first + n - 1
    hf = PatternFill('solid', fgColor='1F4E78'); hfont = Font(color='FFFFFF', bold=True)
    kf = PatternFill('solid', fgColor='FFF2CC')
    good = PatternFill('solid', fgColor='C6EFCE'); bad = PatternFill('solid', fgColor='FFC7CE')
    grey = PatternFill('solid', fgColor='D9D9D9')
    ctr = Alignment(horizontal='center')
    thin = Side(style='thin', color='BFBFBF'); bd = Border(thin, thin, thin, thin)
    wb = Workbook()

    # Marking
    ws = wb.active; ws.title = 'Marking'
    ws.cell(1, 1, 'Student (Sheet/Page)')
    for q in range(1, nq + 1): ws.cell(1, 1 + q, 'Q%d' % q)
    SC = nq + 2; PC = SC + 1; BL = PC + 1; FL = BL + 1
    ws.cell(1, SC, 'Score /%d' % nq); ws.cell(1, PC, 'Percent')
    ws.cell(1, BL, 'Blank'); ws.cell(1, FL, 'Review Qs (verify)')
    for c in range(1, FL + 1):
        x = ws.cell(1, c); x.fill = hf; x.font = hfont; x.alignment = ctr; x.border = bd
    ws.cell(2, 1, 'ANSWER KEY').font = Font(bold=True); ws.cell(2, 1).fill = kf
    for q in range(1, nq + 1):
        c = ws.cell(2, 1 + q, key[q - 1]); c.fill = kf; c.alignment = ctr; c.font = Font(bold=True); c.border = bd
    lastcol = get_column_letter(1 + nq)
    for i, pg in enumerate(pages):
        r = first + i
        ws.cell(r, 1, pg).alignment = ctr
        for q in range(1, nq + 1):
            a = ans[pg][q - 1]
            c = ws.cell(r, 1 + q, a); c.alignment = ctr; c.border = bd
            c.fill = grey if a == '-' else (good if a == key[q - 1] else bad)
        rng = 'B%d:%s%d' % (r, lastcol, r)
        ws.cell(r, SC, '=SUMPRODUCT(--(%s=$B$2:$%s$2))' % (rng, lastcol)).alignment = ctr
        scl = get_column_letter(SC)
        ws.cell(r, PC, '=%s%d/%d' % (scl, r, nq)).number_format = '0.0%'
        ws.cell(r, BL, '=COUNTIF(%s,"-")' % rng).alignment = ctr
        ws.cell(r, FL, ', '.join('Q%d' % q for q in flags[pg]))
    ws.freeze_panes = 'B3'
    ws.column_dimensions['A'].width = 18
    for q in range(1, nq + 1): ws.column_dimensions[get_column_letter(1 + q)].width = 4.5
    for cc in (SC, PC, BL): ws.column_dimensions[get_column_letter(cc)].width = 9
    ws.column_dimensions[get_column_letter(FL)].width = 34

    # Item Analysis
    wa = wb.create_sheet('Item Analysis')
    for j, l in enumerate(['Question', 'Correct Answer', '# Correct', '# Wrong', '# Blank',
                           'Difficulty (%correct)', '#A', '#B', '#C', '#D'], 1):
        c = wa.cell(1, j, l); c.fill = hf; c.font = hfont; c.alignment = ctr; c.border = bd
    for q in range(1, nq + 1):
        r = q + 1; col = get_column_letter(1 + q)
        rng = "Marking!%s%d:%s%d" % (col, first, col, last)
        wa.cell(r, 1, 'Q%d' % q).alignment = ctr
        wa.cell(r, 2, '=Marking!%s2' % col).alignment = ctr
        wa.cell(r, 3, '=COUNTIF(%s,Marking!%s2)' % (rng, col)).alignment = ctr
        wa.cell(r, 4, '=%d-C%d-E%d' % (n, r, r)).alignment = ctr
        wa.cell(r, 5, '=COUNTIF(%s,"-")' % rng).alignment = ctr
        wa.cell(r, 6, '=C%d/%d' % (r, n)).number_format = '0.0%'
        for k, opt in enumerate('ABCD'):
            wa.cell(r, 7 + k, '=COUNTIF(%s,"%s")' % (rng, opt)).alignment = ctr
    base = nq + 2
    wa.cell(base + 1, 1, 'Hardest Q (lowest % correct)').font = Font(bold=True)
    wa.cell(base + 1, 3, '=INDEX(A2:A%d,MATCH(MIN(F2:F%d),F2:F%d,0))' % (nq + 1, nq + 1, nq + 1))
    wa.cell(base + 2, 1, 'Easiest Q (highest % correct)').font = Font(bold=True)
    wa.cell(base + 2, 3, '=INDEX(A2:A%d,MATCH(MAX(F2:F%d),F2:F%d,0))' % (nq + 1, nq + 1, nq + 1))
    wa.cell(base + 3, 1, 'Average difficulty').font = Font(bold=True)
    wa.cell(base + 3, 3, '=AVERAGE(F2:F%d)' % (nq + 1)).number_format = '0.0%'
    for col, w in zip('ABCDEFGHIJ', [10, 14, 10, 9, 9, 20, 6, 6, 6, 6]):
        wa.column_dimensions[col].width = w
    wa.freeze_panes = 'A2'

    # Summary
    ws3 = wb.create_sheet('Summary', 0)
    scl = get_column_letter(SC); pcl = get_column_letter(PC)
    S = 'Marking!%s%d:%s%d' % (scl, first, scl, last)
    P = 'Marking!%s%d:%s%d' % (pcl, first, pcl, last)
    ws3.cell(1, 1, 'MC Marking - Class Summary').font = Font(bold=True, size=14)
    rows = [('Number of students', '=COUNT(%s)' % S, '0'),
            ('Total marks (each)', str(nq), '0'),
            ('Average score', '=AVERAGE(%s)' % S, '0.00'),
            ('Average percent', '=AVERAGE(%s)' % P, '0.0%'),
            ('Median score', '=MEDIAN(%s)' % S, '0.0'),
            ('Highest score', '=MAX(%s)' % S, '0'),
            ('Lowest score', '=MIN(%s)' % S, '0'),
            ('Std deviation', '=STDEV(%s)' % S, '0.00')]
    if pass_mark is not None:
        rows += [('Pass count (>=%d%%)' % round(pass_mark * 100), '=COUNTIF(%s,">=%s")' % (P, pass_mark), '0'),
                 ('Pass rate', '=COUNTIF(%s,">=%s")/COUNT(%s)' % (P, pass_mark, P), '0.0%')]
    r = 3
    for lab, f, fmt in rows:
        ws3.cell(r, 1, lab).font = Font(bold=True)
        ws3.cell(r, 2, f).number_format = fmt; r += 1
    ws3.column_dimensions['A'].width = 24; ws3.column_dimensions['B'].width = 12

    # Notes
    wn = wb.create_sheet('Notes')
    for i, t in enumerate([
        'Generated by mark_mc.py',
        '- One scanned page = one student; bubbles read by darkness, darkest box per Q wins.',
        '- Green = correct, red = wrong, grey = blank. KEY row highlighted yellow.',
        '- "Review Qs" = cells auto-flagged as blank / faint / double / crossed-out and checked by hand.',
        '- "X" in a cell = invalid double-mark (counts wrong).',
        '- Score, Percent, Item Analysis and Summary are live Excel formulas; edit a cell and they update.',
        '- Student identifier = scan page number.'], 1):
        wn.cell(i, 1, t)
    wn.column_dimensions['A'].width = 110
    wb.save(out)


def main():
    key = load_key(KEY)
    nq = NUM_Q or len(key)
    key = key[:nq]
    g = GEOMETRY
    T = targets(nq, g)

    if '--calibrate' in sys.argv:
        calibrate(PDF_PATH, DPI, nq, g)
        return

    work = "_pages"
    paths = render_pages(PDF_PATH, DPI, work)
    ans, flags, detail = {}, {}, {}
    for p in paths:
        pg = int(os.path.basename(p).split('_')[1].split('.')[0])
        res = read_page(p, T, g)
        detail[pg] = res
        row, fl = [], set()
        for q in range(1, nq + 1):
            a = res[q]['ans'] or '-'
            if res[q]['flag']: fl.add(q)
            if (pg, q) in OVERRIDES:
                a = OVERRIDES[(pg, q)]; fl.add(q)
            row.append(a)
        ans[pg] = row; flags[pg] = sorted(fl)
    pages = sorted(ans)
    json.dump(detail, open('detect.json', 'w'))
    build_excel(pages, ans, flags, key, OUTPUT, PASS_MARK)

    # cleanup intermediate renders
    for p in paths: os.remove(p)
    try: os.rmdir(work)
    except OSError: pass

    sc = [sum(1 for q in range(nq) if ans[pg][q] == key[q]) for pg in pages]
    print("Marked %d sheets, %d questions -> %s" % (len(pages), nq, OUTPUT))
    print("Class average: %.2f / %d (%.1f%%)" % (sum(sc) / len(sc), nq, sum(sc) / len(sc) / nq * 100))
    print("Flagged cells to review: %d" % sum(len(f) for f in flags.values()))


if __name__ == '__main__':
    main()
